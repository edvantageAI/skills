#!/usr/bin/env python3
"""
Create document with transcript-screenshot table in various formats.
Supports: HTML, DOCX, and Google Docs (via API).
"""

import json
import argparse
import sys
import os
from pathlib import Path
import base64
import yaml
import logging

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


def load_config(config_path=None):
    """Load configuration from YAML file."""
    if config_path is None:
        config_path = Path(__file__).parent.parent / 'config.yaml'

    config_path = Path(config_path)
    if config_path.exists():
        try:
            with open(config_path, 'r') as f:
                return yaml.safe_load(f)
        except Exception as e:
            logger.warning(f"Failed to load config from {config_path}: {e}")
            return {}
    return {}


def create_html_document(aligned_data, output_path):
    """Create HTML document with embedded images."""
    html = """<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>Video Lecture Notes</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            margin: 20px;
            background-color: #f5f5f5;
        }
        h1 {
            color: #333;
            border-bottom: 2px solid #4285f4;
            padding-bottom: 10px;
        }
        .info {
            background-color: #e8f0fe;
            padding: 10px;
            border-radius: 5px;
            margin-bottom: 20px;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            background-color: white;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        th {
            background-color: #4285f4;
            color: white;
            padding: 12px;
            text-align: left;
            font-weight: bold;
        }
        td {
            padding: 12px;
            border-bottom: 1px solid #ddd;
            vertical-align: top;
        }
        tr:hover {
            background-color: #f5f5f5;
        }
        .timestamp {
            color: #666;
            font-size: 0.9em;
            font-style: italic;
        }
        .screenshot {
            max-width: 100%;
            height: auto;
            border: 1px solid #ddd;
            border-radius: 4px;
        }
        .transcript-cell {
            width: 50%;
        }
        .screenshot-cell {
            width: 50%;
            text-align: center;
        }
    </style>
</head>
<body>
    <h1>Video Lecture Notes</h1>
    <div class="info">
        <p><strong>Video:</strong> {video_path}</p>
        <p><strong>Duration:</strong> {duration:.2f} seconds ({duration_min:.1f} minutes)</p>
        <p><strong>Total Keyframes:</strong> {total_keyframes}</p>
    </div>
    <table>
        <thead>
            <tr>
                <th>Transcript</th>
                <th>Screenshot</th>
            </tr>
        </thead>
        <tbody>
"""
    
    video_info = aligned_data['video_info']
    aligned_items = aligned_data['aligned_items']
    
    html = html.format(
        video_path=video_info.get('video_path', 'N/A'),
        duration=video_info.get('duration', 0),
        duration_min=video_info.get('duration', 0) / 60,
        total_keyframes=len(aligned_items)
    )
    
    for item in aligned_items:
        keyframe = item['keyframe']
        transcript = item['transcript']
        
        # Read and encode image as base64
        img_path = keyframe['filepath']
        if os.path.exists(img_path):
            with open(img_path, 'rb') as img_file:
                img_data = base64.b64encode(img_file.read()).decode('utf-8')
                img_src = f"data:image/jpeg;base64,{img_data}"
        else:
            img_src = ""
        
        timestamp = keyframe['timestamp']
        minutes = int(timestamp // 60)
        seconds = int(timestamp % 60)
        
        html += f"""
            <tr>
                <td class="transcript-cell">
                    <div class="timestamp">⏱️ {minutes}:{seconds:02d} (Frame {keyframe['frame_number']})</div>
                    <p>{transcript if transcript else '<em>No transcript available</em>'}</p>
                </td>
                <td class="screenshot-cell">
                    <img src="{img_src}" alt="Screenshot at {timestamp}s" class="screenshot">
                </td>
            </tr>
"""
    
    html += """
        </tbody>
    </table>
</body>
</html>
"""
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

    logger.info(f"✅ HTML document created: {output_path}")


def create_docx_document(aligned_data, output_path, image_width_inches=3.0):
    """Create DOCX document with table."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.error("python-docx not installed. Install with: pip install python-docx")
        sys.exit(1)
    
    doc = Document()
    
    # Add title
    title = doc.add_heading('Video Lecture Notes', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add video info
    video_info = aligned_data['video_info']
    info_para = doc.add_paragraph()
    info_para.add_run('Video: ').bold = True
    info_para.add_run(f"{video_info.get('video_path', 'N/A')}\n")
    info_para.add_run('Duration: ').bold = True
    duration = video_info.get('duration', 0)
    info_para.add_run(f"{duration:.2f} seconds ({duration/60:.1f} minutes)\n")
    info_para.add_run('Total Keyframes: ').bold = True
    info_para.add_run(f"{len(aligned_data['aligned_items'])}\n")
    
    doc.add_paragraph()  # Spacing
    
    # Create table
    aligned_items = aligned_data['aligned_items']
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    
    # Header row
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Transcript'
    header_cells[1].text = 'Screenshot'
    
    # Make header bold
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(12)
    
    # Add data rows
    for item in aligned_items:
        keyframe = item['keyframe']
        transcript = item['transcript']
        
        row_cells = table.add_row().cells
        
        # Transcript cell
        timestamp = keyframe['timestamp']
        minutes = int(timestamp // 60)
        seconds = int(timestamp % 60)
        
        transcript_para = row_cells[0].paragraphs[0]
        timestamp_run = transcript_para.add_run(f"⏱️ {minutes}:{seconds:02d} (Frame {keyframe['frame_number']})\n")
        timestamp_run.font.italic = True
        timestamp_run.font.size = Pt(9)
        timestamp_run.font.color.rgb = RGBColor(102, 102, 102)
        
        text_run = transcript_para.add_run(transcript if transcript else 'No transcript available')
        text_run.font.size = Pt(11)
        
        # Screenshot cell
        img_path = keyframe['filepath']
        if os.path.exists(img_path):
            screenshot_para = row_cells[1].paragraphs[0]
            screenshot_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = screenshot_para.add_run()
            run.add_picture(img_path, width=Inches(image_width_inches))

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(str(output_path))
    logger.info(f"✅ DOCX document created: {output_path}")

    return output_path


def create_google_doc(aligned_data, output_path, credentials_path=None, image_width_inches=3.0):
    """Create Google Doc by uploading DOCX to Drive and converting to Google Docs format."""
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload
    except ImportError:
        logger.error("Google API libraries not installed")
        logger.error("Install with: pip install google-auth google-auth-oauthlib google-auth-httplib2 google-api-python-client")
        sys.exit(1)

    if not credentials_path or not Path(credentials_path).exists():
        logger.error("Google API credentials file not found")
        logger.error("Please provide valid credentials JSON file path with --credentials")
        sys.exit(1)

    # Authenticate
    SCOPES = ['https://www.googleapis.com/auth/drive.file']

    try:
        creds = service_account.Credentials.from_service_account_file(
            credentials_path, scopes=SCOPES)
    except Exception as e:
        logger.error(f"Error loading credentials: {e}", exc_info=True)
        sys.exit(1)

    # Step 1: Create DOCX in temp file
    import tempfile

    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        tmp_docx_path = tmp.name

    try:
        logger.info("Creating temporary DOCX file...")
        create_docx_document(aligned_data, tmp_docx_path, image_width_inches)

        # Step 2: Upload to Google Drive and convert to Docs
        logger.info("Uploading to Google Drive...")
        drive_service = build('drive', 'v3', credentials=creds)

        file_metadata = {
            'name': 'Video Lecture Notes',
            'mimeType': 'application/vnd.google-apps.document'  # Convert to Google Docs format
        }

        media = MediaFileUpload(
            tmp_docx_path,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            resumable=True
        )

        file = drive_service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id,webViewLink'
        ).execute()

        doc_id = file.get('id')
        doc_url = file.get('webViewLink')

        # Save info to output file
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        with open(output_path, 'w') as f:
            json.dump({
                'doc_id': doc_id,
                'doc_url': doc_url,
                'method': 'docx_upload',
                'created_at': str(Path(tmp_docx_path).stat().st_mtime)
            }, f, indent=2)

        logger.info(f"✅ Google Doc created: {output_path}")
        logger.info(f"   View at: {doc_url}")

    finally:
        # Clean up temp file
        Path(tmp_docx_path).unlink(missing_ok=True)


def main():
    # Load configuration
    config = load_config()
    doc_config = config.get('document_generation', {})

    # Set defaults from config
    default_format = doc_config.get('default_format', 'html')
    image_width = doc_config.get('docx', {}).get('image_width_inches', 3.0)

    parser = argparse.ArgumentParser(description='Create document from aligned transcript and keyframes')
    parser.add_argument('aligned_json', help='Path to aligned JSON file')
    parser.add_argument('output_path', help='Path to output file')
    parser.add_argument('--format', choices=['html', 'docx', 'gdoc'], default=default_format,
                        help=f'Output format (default: {default_format})')
    parser.add_argument('--credentials', help='Path to Google API credentials JSON (required for gdoc format)')
    parser.add_argument('--image-width', type=float, default=image_width,
                        help=f'Image width in inches for DOCX/GDoc (default: {image_width})')
    parser.add_argument('--config', help='Path to config file')

    args = parser.parse_args()

    # Reload config if custom path provided
    if args.config:
        config = load_config(args.config)
        doc_config = config.get('document_generation', {})

    # Load aligned data
    try:
        aligned_json_path = Path(args.aligned_json)
        if not aligned_json_path.exists():
            raise FileNotFoundError(f"Aligned JSON file not found: {aligned_json_path}")

        with open(aligned_json_path, 'r') as f:
            aligned_data = json.load(f)

        if 'aligned_items' not in aligned_data:
            raise ValueError("Invalid aligned data: missing 'aligned_items' field")

        logger.info(f"Loaded aligned data with {len(aligned_data['aligned_items'])} items")
    except Exception as e:
        logger.error(f"Error loading aligned data: {e}", exc_info=True)
        sys.exit(1)

    # Create document based on format
    try:
        if args.format == 'html':
            create_html_document(aligned_data, args.output_path)
        elif args.format == 'docx':
            create_docx_document(aligned_data, args.output_path, image_width_inches=args.image_width)
        elif args.format == 'gdoc':
            create_google_doc(aligned_data, args.output_path, args.credentials, image_width_inches=args.image_width)
        else:
            logger.error(f"Unknown format: {args.format}")
            sys.exit(1)
    except Exception as e:
        logger.error(f"Error creating document: {e}", exc_info=True)
        sys.exit(1)


if __name__ == '__main__':
    main()
